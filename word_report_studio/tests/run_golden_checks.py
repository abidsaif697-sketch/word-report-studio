#!/usr/bin/env python3
"""
run_golden_checks.py — full end-to-end pipeline checks, not just unit-level
function calls (see run_regression.py for those). Every check here freezes
a bug that was found by actually RUNNING the app and looking at real output
(rendered .docx/.pdf pages) rather than by reading code or unit-testing a
function in isolation. If a future change reintroduces one of these, this
suite catches it before a person has to notice a garbled page again.

Some checks need LibreOffice (PDF export); they SKIP with a note instead of
failing when it isn't installed, matching the rest of the project's offline-
degrades-gracefully philosophy.

Usage:  python tests/run_golden_checks.py
"""
from __future__ import annotations
import json
import os
import subprocess
import sys
import shutil
import tempfile
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.chdir(ROOT)

PASS, FAIL, SKIP = 0, 0, 0


def check(name, cond, detail=""):
    global PASS, FAIL
    ok = bool(cond)
    PASS += ok
    FAIL += (not ok)
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not ok else ""))


def skip(name, reason):
    global SKIP
    SKIP += 1
    print(f"  SKIP  {name}  ({reason})")


def section(title):
    print(f"\n== {title} ==")


def page_break_count(docx_path) -> int:
    from docx import Document
    d = Document(docx_path)
    return sum(1 for p in d.paragraphs if p.paragraph_format.page_break_before)


def build_fixture_many_headings(path, n=60):
    from docx import Document
    d = Document()
    d.add_heading("Information Security Policy / سياسة أمن المعلومات", level=0)
    for i in range(1, n + 1):
        d.add_heading(f"{i}. Clause {i} requirement / البند {i}", level=1)
        d.add_paragraph(f"All staff must comply with clause {i}. "
                        f"يجب على جميع الموظفين الالتزام بالبند {i}.")
    d.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="wrs_golden_")
    try:
        # ------------------------------------------------ template config invariants
        section("Template config invariants")
        from app.template_manager import TemplateManager
        tm = TemplateManager(os.path.join(ROOT, "templates"))
        template_ids = [t.id for t in tm.list_templates()]
        for tid in template_ids:
            path = os.path.join(ROOT, "templates", tid, "metadata.json")
            if not os.path.isfile(path):
                continue
            with open(path, encoding="utf-8") as f:
                raw = json.load(f)
            if "divider_style" in raw or "opener_motif" in raw:
                check(f"{tid}: section_dividers set (divider config present)",
                      raw.get("section_dividers") is True,
                      f"divider_style/opener_motif configured but section_dividers={raw.get('section_dividers')!r}")

        # ------------------------------------------------ cross-template consistency
        section("Poster divider consistency across templates")
        from app.content_parser import ContentParser
        from app.structure_model import DocumentMeta
        from app.layout_engine import LayoutEngine
        from app.docx_renderer import DocxRenderer

        content = open("sample_input/rough_prose.txt", encoding="utf-8").read()
        break_counts = {}
        premium_ids = [tid for tid in template_ids
                      if json.load(open(os.path.join(ROOT, "templates", tid, "metadata.json"),
                                        encoding="utf-8")).get("section_dividers")]
        for tid in premium_ids:
            parser = ContentParser()
            report = parser.parse_auto(content, meta=DocumentMeta(title="Golden"))
            engine = LayoutEngine(tm)
            options = engine.generate_options(report, max_options=1, template_ids=[tid])
            out = DocxRenderer().render(report, options[0].template, options[0].variant,
                                        os.path.join(tmp, f"consistency_{tid}.docx"))
            break_counts[tid] = page_break_count(out)
        check("all divider-enabled templates produce the same break count for identical input",
              len(set(break_counts.values())) <= 1, str(break_counts))

        # ------------------------------------------------ page-explosion guard
        section("Page-explosion guard (many-headings document)")
        many_path = os.path.join(tmp, "many_headings.docx")
        build_fixture_many_headings(many_path, n=60)
        from app.docx_ingest import docx_to_markdown
        many_content = docx_to_markdown(many_path)
        parser = ContentParser()
        report = parser.parse_auto(many_content, meta=DocumentMeta(title="Many"))
        check("60-heading document parsed as many sections", len(report.sections) >= 50,
              str(len(report.sections)))
        engine = LayoutEngine(tm)
        options = engine.generate_options(report, max_options=1,
                                          template_ids=[premium_ids[0]] if premium_ids else None)
        out = DocxRenderer().render(report, options[0].template, options[0].variant,
                                    os.path.join(tmp, "many_headings_out.docx"))
        breaks = page_break_count(out)
        check("many-heading document does not get a forced break per section",
              breaks <= 6, f"{breaks} forced breaks for {len(report.sections)} sections")

        # ------------------------------------------------ KPI label clause-scoping
        section("KPI label clause-scoping")
        from app.auto_enrich import _kpi_from_text
        text = ("Incident response time averaged 4.2 hours, a 12.4M SAR budget "
                "was allocated, and 3 of 5 regional offices achieved full certification.")
        kpi = _kpi_from_text(text)
        check("12.4M KPI extracted", kpi is not None and kpi.value == "12.4M",
              str(kpi))
        if kpi:
            check("label does not bleed the neighboring 'incident response' clause",
                  "incident response" not in kpi.label.lower(), kpi.label)
        text_ar = ("بلغ معدل إتمام التدريب 87%، وتم تخصيص ميزانية قدرها 12.4 مليون ريال، "
                  "وحقق 3 من 5 مكاتب إقليمية الاعتماد الكامل.")
        kpi_ar = _kpi_from_text(text_ar)
        check("Arabic KPI extracted with correctly scoped label",
              kpi_ar is not None and kpi_ar.value == "87%"
              and "مليون" not in kpi_ar.label, str(kpi_ar))

        # ------------------------------------------------ LLM slot-mapping guards
        section("LLM slot-mapping safety nets (no live model needed)")
        from app.llm_brain import _looks_like_inventory_echo, _slots_json_schema
        from app.template_filler import Slot
        echoed = {1: "fine", 2: "page 1 | micro | max 40 chars | now: leaked prompt text"}
        clean = {1: "Al Riyadh Logistics", 2: "Moving the Kingdom Forward"}
        check("inventory-echo garbage is detected", _looks_like_inventory_echo(echoed))
        check("clean mapping is not flagged", not _looks_like_inventory_echo(clean))
        schema = _slots_json_schema([Slot(slot_id=5, page=1, text="x", capacity=10, kind="micro")])
        check("slot schema locks down keys (additionalProperties False)",
              schema.get("additionalProperties") is False and "5" in schema.get("properties", {}))

        # ------------------------------------------------ PDF export path handling
        section("PDF export (needs LibreOffice)")
        from app import pdf_exporter
        if not pdf_exporter.is_available():
            skip("mixed-separator path export", "LibreOffice not installed")
        else:
            # forward-slash output_dir + os.path.join reproduces the exact
            # mixed-separator path ("out/dir\\file.docx") that broke soffice.
            mixed_dir = os.path.join(tmp, "pdf_check").replace("\\", "/")
            os.makedirs(mixed_dir, exist_ok=True)
            src = os.path.join(tmp, "consistency_" + premium_ids[0] + ".docx") if premium_ids else None
            if src and os.path.isfile(src):
                shutil.copy(src, os.path.join(mixed_dir, "doc.docx"))
                doc_path = mixed_dir + "/doc.docx"  # forward slashes, like a user's --output-dir
                try:
                    pdf = pdf_exporter.convert_to_pdf(doc_path, output_dir=mixed_dir)
                    check("mixed-separator path converts to PDF", os.path.exists(pdf))
                except Exception as e:
                    check("mixed-separator path converts to PDF", False, str(e))
            else:
                skip("mixed-separator path export", "no source docx available")

        # ------------------------------------------------ CLI smoke test (--pdf)
        section("CLI smoke test")
        if not pdf_exporter.is_available():
            skip("run_cli.py --pdf exits 0", "LibreOffice not installed")
        else:
            cli_out = os.path.join(tmp, "cli_out")
            r = subprocess.run(
                [sys.executable, "run_cli.py", "--input", "sample_input/source_policy.docx",
                 "--title", "Golden CLI Check", "--templates", "government_formal",
                 "--num-options", "1", "--output-dir", cli_out, "--pdf", "--ai", "off"],
                cwd=ROOT, capture_output=True, timeout=120, text=True,
            )
            check("run_cli.py --pdf exits 0 (no UnboundLocalError)", r.returncode == 0,
                  r.stderr[-400:] if r.returncode else "")
            produced_pdf = any(f.endswith(".pdf") for f in os.listdir(cli_out)) if os.path.isdir(cli_out) else False
            check("run_cli.py --pdf actually produces a .pdf file", produced_pdf,
                  "PDF export failed silently or was skipped")

    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print(f"\n{'=' * 40}\nRESULT: {PASS} passed, {FAIL} failed, {SKIP} skipped")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
