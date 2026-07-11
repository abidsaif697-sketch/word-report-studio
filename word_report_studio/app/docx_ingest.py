"""
docx_ingest.py
---------------
Read an existing Word document and convert it to the engine's markdown
dialect — the front door for the user's real workflow, where source content
lives in .docx files. Pasting text out of Word destroys tables and bidi
ordering; reading the file directly keeps headings, tables, lists, and
embedded images intact, so the redesign starts from real structure.

The output is markdown text (not a ReportDocument) on purpose: the GUI puts
it in the editor so the user can see and adjust what was extracted before
generating, and the normal parse + auto-enrich pipeline runs on it
unchanged (numeric tables from the source will get charts automatically).
"""

from __future__ import annotations
import os
import re
from typing import List, Optional

from docx import Document
from docx.oxml.ns import qn


_HEADING_STYLE_RE = re.compile(r"heading\s*(\d)", re.IGNORECASE)
# paragraphs to drop: the source's own TOC (we regenerate one) and empties
_SKIP_STYLE_RE = re.compile(r"^(TOC|toc \d|table of figures)", re.IGNORECASE)


def _heading_level(paragraph) -> Optional[int]:
    try:
        style = paragraph.style
        for probe in (getattr(style, "style_id", "") or "",
                      getattr(style, "name", "") or ""):
            m = _HEADING_STYLE_RE.search(probe)
            if m:
                return max(1, min(int(m.group(1)), 4))
        # fall back to explicit outline level
        out = paragraph._p.find(qn("w:pPr") + "/" + qn("w:outlineLvl"))
        if out is not None:
            return max(1, min(int(out.get(qn("w:val"))) + 1, 4))
    except Exception:
        pass
    return None


def _is_list_item(paragraph) -> bool:
    if paragraph._p.pPr is not None and paragraph._p.pPr.find(qn("w:numPr")) is not None:
        return True
    name = (getattr(paragraph.style, "name", "") or "").lower()
    return name.startswith(("list", "قائمة"))


def _is_ordered(paragraph) -> bool:
    name = (getattr(paragraph.style, "name", "") or "").lower()
    return "number" in name


def _cell_text(cell) -> str:
    parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    return " ".join(parts).replace("|", "/")


def _extract_images(paragraph, doc, image_dir: str, counter: List[int]) -> List[str]:
    """Save embedded images to image_dir; return markdown image lines."""
    lines = []
    for blip in paragraph._p.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            part = doc.part.related_parts[rid]
            ext = os.path.splitext(part.partname)[1] or ".png"
            counter[0] += 1
            fname = f"source_image_{counter[0]:03d}{ext}"
            os.makedirs(image_dir, exist_ok=True)
            path = os.path.join(image_dir, fname)
            with open(path, "wb") as f:
                f.write(part.blob)
            lines.append(f"![]({path})")
        except Exception:
            continue
    return lines


def docx_to_markdown(path: str, image_dir: Optional[str] = None) -> str:
    """Convert a source .docx to the engine's markdown dialect."""
    doc = Document(path)
    image_dir = image_dir or os.path.join(
        os.path.dirname(os.path.abspath(path)),
        "_extracted_images_" + re.sub(r"\W+", "_", os.path.splitext(os.path.basename(path))[0]))
    out: List[str] = []
    img_counter = [0]
    pending_list: List[str] = []
    pending_ordered = False

    def flush_list():
        nonlocal pending_list, pending_ordered
        for i, item in enumerate(pending_list, start=1):
            out.append(f"{i}. {item}" if pending_ordered else f"- {item}")
        if pending_list:
            out.append("")
        pending_list = []

    # walk body children in document order so tables land where they belong
    body = doc.element.body
    par_map = {p._p: p for p in doc.paragraphs}
    tbl_map = {t._tbl: t for t in doc.tables}

    for child in body:
        if child.tag == qn("w:p"):
            p = par_map.get(child)
            if p is None:
                continue
            style_name = getattr(p.style, "name", "") or ""
            if _SKIP_STYLE_RE.match(style_name):
                continue
            text = p.text.strip()
            imgs = _extract_images(p, doc, image_dir, img_counter)
            level = _heading_level(p)
            if level is not None and text:
                flush_list()
                out.append("#" * level + " " + text)
                out.append("")
            elif _is_list_item(p) and text:
                if pending_list and pending_ordered != _is_ordered(p):
                    flush_list()
                pending_ordered = _is_ordered(p)
                pending_list.append(text)
            elif text:
                flush_list()
                out.append(text)
                out.append("")
            for line in imgs:
                flush_list()
                out.append(line)
                out.append("")
        elif child.tag == qn("w:tbl"):
            flush_list()
            t = tbl_map.get(child)
            if t is None or not t.rows:
                continue
            n_cols = max(len(r.cells) for r in t.rows)
            for ri, row in enumerate(t.rows):
                cells = [_cell_text(c) for c in row.cells]
                cells += [""] * (n_cols - len(cells))
                out.append("| " + " | ".join(cells) + " |")
                if ri == 0:
                    out.append("|" + "---|" * n_cols)
            out.append("")
    flush_list()

    md = "\n".join(out)
    md = re.sub(r"\n{3,}", "\n\n", md).strip() + "\n"
    return md


def source_title(path: str) -> Optional[str]:
    """Title from the source document's core properties, if set."""
    try:
        title = Document(path).core_properties.title
        return title.strip() if title and title.strip() else None
    except Exception:
        return None
