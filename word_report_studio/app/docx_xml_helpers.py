"""
docx_xml_helpers.py
---------------------
Low-level OOXML helpers that python-docx does not expose through its public
API: RTL/bidi flags (needed for correct Arabic shaping and reading order),
complex-script fonts, paragraph shading (for colored bands/callout boxes),
page-number fields, and a TOC field.

These are standard, widely-used OOXML recipes (python-docx only wraps a
subset of the ECMA-376 spec; anything to do with bidi/complex-script runs,
field codes, and cell/paragraph shading has to be added by hand via oxml).
"""

from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor


def set_paragraph_bidi(paragraph, rtl: bool = True):
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:bidi"))
    if rtl:
        if existing is None:
            pPr.append(OxmlElement("w:bidi"))
    else:
        if existing is not None:
            pPr.remove(existing)


def set_run_rtl(run, rtl: bool = True):
    rPr = run._r.get_or_add_rPr()
    existing = rPr.find(qn("w:rtl"))
    if rtl:
        if existing is None:
            rPr.append(OxmlElement("w:rtl"))
    else:
        if existing is not None:
            rPr.remove(existing)


def set_run_font(run, ascii_font=None, cs_font=None, size_pt=None,
                  bold=None, italic=None, underline=None, color_hex=None):
    """Set Latin (ascii/hAnsi) font via python-docx's normal API, and the
    complex-script (cs) font via raw XML so Arabic renders with the right
    typeface even inside a Latin-styled paragraph."""
    if ascii_font:
        run.font.name = ascii_font
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if underline is not None:
        run.font.underline = underline
    if color_hex:
        try:
            run.font.color.rgb = RGBColor.from_string(color_hex.replace("#", ""))
        except Exception:
            pass

    rPr = run._r.get_or_add_rPr()
    if cs_font:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:cs"), cs_font)
    if size_pt:
        szCs = rPr.find(qn("w:szCs"))
        if szCs is None:
            szCs = OxmlElement("w:szCs")
            rPr.append(szCs)
        szCs.set(qn("w:val"), str(int(round(size_pt * 2))))


def set_paragraph_shading(paragraph, hex_color: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.replace("#", ""))
    pPr.append(shd)


def set_cell_background(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.replace("#", ""))
    tcPr.append(shd)


def set_paragraph_border_bottom(paragraph, hex_color: str, size: int = 24):
    """Adds a colored horizontal rule under a paragraph (used for minimal_top cover)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), hex_color.replace("#", ""))
    pBdr.append(bottom)


def set_paragraph_border_left(paragraph, hex_color: str, size: int = 24, space: int = 8):
    """Colored vertical rule to the left of a paragraph (pull-quote look)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), hex_color.replace("#", ""))
    pBdr.append(left)


def set_mirror_margins(section, mirror: bool = True):
    try:
        sectPr = section._sectPr
        existing = sectPr.find(qn("w:mirrorMargins"))
        if mirror:
            if existing is None:
                sectPr.append(OxmlElement("w:mirrorMargins"))
        else:
            if existing is not None:
                sectPr.remove(existing)
    except Exception:
        pass  # cosmetic only; never fail the render because of this


def add_page_number_field(paragraph, prefix: str = ""):
    if prefix:
        r = paragraph.add_run(prefix)
    fldChar_begin = OxmlElement("w:fldChar")
    fldChar_begin.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar_end = OxmlElement("w:fldChar")
    fldChar_end.set(qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    run._r.append(fldChar_begin)
    run._r.append(instrText)
    run._r.append(fldChar_end)


def add_toc_field(document, heading_levels: str = "1-4",
                   placeholder_text: str = "Right-click here and choose \"Update Field\" "
                                            "(or press F9) to generate the Table of Contents."):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = f'TOC \\o "{heading_levels}" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = placeholder_text
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")

    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(placeholder)
    r.append(fldChar3)
    return paragraph


def set_table_cell_margins(table, top_dxa: int = 100, bottom_dxa: int = 100,
                            left_dxa: int = 115, right_dxa: int = 115):
    """Default interior padding for every cell in a table (1440 dxa = 1 inch).
    This is most of what makes a table look 'designed' instead of cramped."""
    tblPr = table._tbl.tblPr
    existing = tblPr.find(qn("w:tblCellMar"))
    if existing is not None:
        tblPr.remove(existing)
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top_dxa), ("bottom", bottom_dxa),
                     ("start", left_dxa), ("end", right_dxa)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_table_width(table, width_dxa: int, indent_dxa: int = 0):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(width_dxa))
    tblW.set(qn("w:type"), "dxa")

    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent_dxa))
    tblInd.set(qn("w:type"), "dxa")


def set_table_fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_table_grid(table, widths_dxa):
    tblGrid = table._tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        table._tbl.remove(tblGrid)
    tblGrid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width)))
        tblGrid.append(col)
    table._tbl.insert(0, tblGrid)


def set_cell_width(cell, width_dxa: int):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_dxa)))
    tcW.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    existing = trPr.find(qn("w:tblHeader"))
    if existing is None:
        trPr.append(OxmlElement("w:tblHeader"))


def set_row_cant_split(row):
    """Forbid Word from breaking a table row across pages — used on one-row
    visuals (callouts, KPI strips, section dividers) so they never leave an
    orphaned fragment on the next page."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def _border_el(tag: str, val: str, sz: int, color: str):
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))          # eighths of a point
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color.replace("#", ""))
    return el


def set_table_borders(table, edges: dict):
    """edges: {"top"|"bottom"|"start"|"end"|"insideH"|"insideV":
               (val, sz_eighths_pt, hex_color) or None to remove}.
    Only the listed edges are touched."""
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for tag, spec in edges.items():
        existing = borders.find(qn(f"w:{tag}"))
        if existing is not None:
            borders.remove(existing)
        if spec is None:
            borders.append(_border_el(tag, "none", 0, "auto"))
        else:
            val, sz, color = spec
            borders.append(_border_el(tag, val, sz, color))


def set_cell_borders(cell, edges: dict):
    """Same spec format as set_table_borders, applied to a single cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for tag, spec in edges.items():
        existing = borders.find(qn(f"w:{tag}"))
        if existing is not None:
            borders.remove(existing)
        if spec is None:
            borders.append(_border_el(tag, "none", 0, "auto"))
        else:
            val, sz, color = spec
            borders.append(_border_el(tag, val, sz, color))


def set_row_height(row, inches: float, rule: str = "atLeast"):
    """Force a table row to a height — how covers get full-page color areas."""
    trPr = row._tr.get_or_add_trPr()
    existing = trPr.find(qn("w:trHeight"))
    if existing is not None:
        trPr.remove(existing)
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(inches * 1440)))
    h.set(qn("w:hRule"), rule)
    trPr.append(h)


def set_vertical_alignment(cell, align: str = "center"):
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:vAlign"))
    if existing is not None:
        tcPr.remove(existing)
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), align)
    tcPr.append(v)


def safe_set_table_style(table, style_name: str, fallback: str = "Table Grid"):
    for name in (style_name, fallback, None):
        try:
            if name:
                table.style = name
            return
        except (KeyError, ValueError):
            continue
